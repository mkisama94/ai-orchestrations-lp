from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph


BASE_DIR = Path(__file__).resolve().parent
SOURCE = BASE_DIR / "source.docx"
OUTPUT = BASE_DIR / "職務経歴書(山本)2026_更新版.docx"


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    """Replace text while preserving paragraph properties and representative run styling."""
    template_run = next((run for run in paragraph.runs if run.text), None)
    if template_run is None and paragraph.runs:
        template_run = paragraph.runs[0]
    template_rpr = deepcopy(template_run._r.get_or_add_rPr()) if template_run else None

    for child in list(paragraph._p):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            paragraph._p.remove(child)

    run = paragraph.add_run(text)
    if template_rpr is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, template_rpr)


def clone_paragraph(template: Paragraph, text: str, parent) -> Paragraph:
    paragraph_xml = deepcopy(template._p)
    paragraph = Paragraph(paragraph_xml, parent)
    set_paragraph_text(paragraph, text)
    return paragraph


def replace_cell_paragraphs(cell: _Cell, paragraphs: list[Paragraph]) -> None:
    for paragraph_xml in list(cell._tc.p_lst):
        cell._tc.remove(paragraph_xml)
    for paragraph in paragraphs:
        cell._tc.append(paragraph._p)


def build_updated_resume() -> Path:
    document = Document(SOURCE)

    # Refresh the "as of" date and make the new independent status explicit in the summary.
    set_paragraph_text(document.paragraphs[1], "2026年8月20日現在")
    set_paragraph_text(
        document.paragraphs[5],
        "2025年11月にコンサルタントとして独立。現在は、事業立ち上げから現場改善までを一気通貫で支援する「End-to-End ITコンサルタント」として活動しています。",
    )

    # Add the latest independent career entry at the top of the detailed work history.
    section_table = document.tables[5]
    table_template = document.tables[6]
    heading_template = document.paragraphs[46]
    business_template = document.paragraphs[48]
    blank_template = document.paragraphs[49]

    heading = clone_paragraph(
        heading_template,
        "独立／個人事業（2025年11月～現在）",
        document._body,
    )
    business = clone_paragraph(
        business_template,
        "事業内容： コンサルティング",
        document._body,
    )
    independent_table_xml = deepcopy(table_template._tbl)
    independent_table = Table(independent_table_xml, document._body)
    trailing_blank = clone_paragraph(blank_template, "", document._body)

    period_templates = table_template.rows[1].cells[0].paragraphs
    replace_cell_paragraphs(
        independent_table.rows[1].cells[0],
        [
            clone_paragraph(period_templates[0], "2025年11月", independent_table.rows[1].cells[0]),
            clone_paragraph(period_templates[1], "　～", independent_table.rows[1].cells[0]),
            clone_paragraph(period_templates[2], "現在", independent_table.rows[1].cells[0]),
        ],
    )

    detail_templates = table_template.rows[1].cells[1].paragraphs
    replace_cell_paragraphs(
        independent_table.rows[1].cells[1],
        [
            clone_paragraph(detail_templates[0], "概要", independent_table.rows[1].cells[1]),
            clone_paragraph(
                detail_templates[1],
                "2025年11月、コンサルタントとして独立。",
                independent_table.rows[1].cells[1],
            ),
            clone_paragraph(
                detail_templates[2],
                "事業立ち上げから現場改善まで、IT・DX領域の支援に従事。",
                independent_table.rows[1].cells[1],
            ),
            clone_paragraph(detail_templates[3], "", independent_table.rows[1].cells[1]),
            clone_paragraph(detail_templates[4], "＜主な実績＞", independent_table.rows[1].cells[1]),
            clone_paragraph(
                detail_templates[5],
                "■上記「独立後 案件実績」参照",
                independent_table.rows[1].cells[1],
            ),
        ],
    )

    org_templates = table_template.rows[1].cells[2].paragraphs
    replace_cell_paragraphs(
        independent_table.rows[1].cells[2],
        [clone_paragraph(org_templates[1], "個人", independent_table.rows[1].cells[2])],
    )

    anchor = section_table._tbl
    for block in (heading._p, business._p, independent_table._tbl, trailing_blank._p):
        anchor.addnext(block)
        anchor = block

    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_updated_resume())
